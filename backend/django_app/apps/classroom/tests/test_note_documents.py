"""Focused regression coverage; shared setup lives in the fixture module."""
from io import BytesIO
from unittest.mock import patch
import zipfile
from django.test import TestCase
from docx import Document
from apps.classroom.services.lecture_notes import save_lecture_note
from .lecture_note_fixtures import LectureNotesFixture, lesson_answer


class NoteDocumentTests(LectureNotesFixture, TestCase):
    @patch('apps.classroom.services.note_translation.get_llm')
    def test_list_and_word_download_reuse_saved_content_without_llm(self, llm):
        response = self.client.get('/api/notes/', {'session_id': 'lesson-session'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data[0]['sections'][0]['formula'], 'V = I × R')
        for _ in range(2):
            result = self.client.get(f'/api/notes/{self.note.id}/document/', {'language': 'en'})
            self.assertEqual(result.status_code, 200)
            document = Document(BytesIO(result.content))
            text = '\n'.join(p.text for p in document.paragraphs)
            self.assertIn('12 V / 4 ohms = 3 A', text)
            self.assertIn('V = I × R', text)
        llm.assert_not_called()


    def test_word_export_formats_markdown_and_preserves_all_words(self):
        answer = lesson_answer('ur')
        answer['chunks'][0]['speak'] = 'The **midbrain** controls *movement*.\n\n1. Eye movement\n2. Hearing'
        note = save_lecture_note(self.user, 'formatted-session', answer)
        result = self.client.get(f'/api/notes/{note.id}/document/', {'language': 'ur'})
        self.assertEqual(result.status_code, 200)
        document = Document(BytesIO(result.content))
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertIn('The midbrain controls movement.', text)
        self.assertIn('Eye movement', text)
        self.assertIn('Hearing', text)
        self.assertNotIn('**', text)
        self.assertTrue(any(run.text == 'midbrain' and run.bold for p in document.paragraphs for run in p.runs))
        self.assertTrue(any(run.text == 'movement' and run.italic for p in document.paragraphs for run in p.runs))
        self.assertEqual(sum(p.style.name == 'List Number' for p in document.paragraphs), 2)
        xml = zipfile.ZipFile(BytesIO(result.content)).read('word/document.xml')
        self.assertIn(b'w:bidi', xml)


    def test_each_topic_download_has_its_own_name_date_and_content(self):
        other_answer = lesson_answer()
        other_answer['topic'] = 'Midbrain'
        other_answer['chunks'][0]['speak'] = 'The midbrain controls visual reflexes.'
        other_note = save_lecture_note(self.user, 'midbrain-session', other_answer)
        response = self.client.get(f'/api/notes/{other_note.id}/document/', {'language': 'en'})
        self.assertEqual(response['Content-Disposition'], 'attachment; filename="Midbrain.docx"')
        document = Document(BytesIO(response.content))
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertEqual(document.paragraphs[0].text, 'Midbrain')
        self.assertIn('Neurolearn · Lecture notes', text)
        self.assertIn(f'Lecture date (UTC): {other_note.created_at:%Y-%m-%d}', text)
        self.assertIn('Language: en', text)
        self.assertIn('visual reflexes', text)
        self.assertNotIn('12 V / 4 ohms', text)


    @patch('apps.classroom.services.note_translation.get_llm')
    def test_download_never_implicitly_translates(self, llm):
        result = self.client.get(f'/api/notes/{self.note.id}/document/', {'language': 'ur'})
        self.assertEqual(result.status_code, 400)
        llm.assert_not_called()


    def test_urdu_document_contains_rtl_paragraph_and_run_properties(self):
        note = save_lecture_note(self.user, 'lesson-session', lesson_answer('ur'))
        response = self.client.get(f'/api/notes/{note.id}/document/', {'language': 'ur'})
        self.assertEqual(response.status_code, 200)
        xml = zipfile.ZipFile(BytesIO(response.content)).read('word/document.xml')
        self.assertIn(b'w:bidi', xml)
        self.assertIn(b'w:rtl', xml)


    def test_regional_source_language_is_normalized_for_original_export(self):
        note = save_lecture_note(self.user, 'regional-session', lesson_answer('en-US'))
        self.assertEqual(note.answer['language'], 'en')
        response = self.client.get(f'/api/notes/{note.id}/document/', {'language': 'en'})
        self.assertEqual(response.status_code, 200)
