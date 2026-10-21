"""Render plain and Markdown lecture text as Word paragraphs and runs."""
from markdown_it import MarkdownIt
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ..constants.notes import NOTE_DOCUMENT_FONT, NOTE_CODE_FONT, NOTE_MAX_HEADING_LEVEL

def add_note_paragraph(document, text, *, rtl=False, style=None, inline_tokens=None):
    paragraph = document.add_paragraph(style=style)
    bold = italic = False
    runs = []
    for token in inline_tokens or []:
        if token.type == 'strong_open':
            bold = True
        elif token.type == 'strong_close':
            bold = False
        elif token.type == 'em_open':
            italic = True
        elif token.type == 'em_close':
            italic = False
        elif token.type in ('text', 'code_inline', 'softbreak', 'hardbreak'):
            run = paragraph.add_run('\n' if token.type.endswith('break') else token.content)
            run.bold, run.italic = bold, italic
            if token.type == 'code_inline':
                run.font.name = NOTE_CODE_FONT
            runs.append(run)
    if inline_tokens is None:
        runs.append(paragraph.add_run(text))
    if rtl:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        bidi = OxmlElement('w:bidi')
        paragraph._p.get_or_add_pPr().append(bidi)
        for run in runs:
            run.font.rtl = True
            run.font.complex_script = True
            run.font.name = NOTE_DOCUMENT_FONT
    return paragraph


def add_note_markdown(document, text, *, rtl=False):
    """Render Markdown as Word paragraphs and styled runs without literal markers."""
    tokens = MarkdownIt('commonmark', {'html': False}).parse(text)
    lists = []
    heading = None
    for token in tokens:
        if token.type in ('ordered_list_open', 'bullet_list_open'):
            lists.append('List Number' if token.type == 'ordered_list_open' else 'List Bullet')
        elif token.type in ('ordered_list_close', 'bullet_list_close'):
            lists.pop()
        elif token.type == 'heading_open':
            heading = f'Heading {min(int(token.tag[1:]), NOTE_MAX_HEADING_LEVEL)}'
        elif token.type == 'heading_close':
            heading = None
        elif token.type == 'inline':
            add_note_paragraph(document, '', rtl=rtl, style=heading or (lists[-1] if lists else None),
                               inline_tokens=token.children)
        elif token.type in ('fence', 'code_block'):
            paragraph = add_note_paragraph(document, token.content.rstrip('\n'))
            for run in paragraph.runs:
                run.font.name = NOTE_CODE_FONT


